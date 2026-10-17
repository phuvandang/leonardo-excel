"""Bộ chuẩn trình bày văn bản Word của Leonardo: lề, phông, hệ tiêu đề, bảng, đầu trang/chân trang, khối ký.

Dùng chung bảng màu, logo và thông tin công ty với phần Excel (thư mục excel/ đặt cạnh). Mọi hằng số trình bày nằm ở đây.
"""
import json, os
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _shared(name):
    """Tìm file dùng chung (config.json, logo): trong thư mục word/ trước, rồi tới thư mục excel/ đặt cạnh."""
    for base in (ROOT, os.path.join(os.path.dirname(ROOT), 'excel')):
        p = os.path.join(base, name)
        if os.path.exists(p): return p
    raise FileNotFoundError(f'Không tìm thấy {name}: thiếu thư mục excel/ của skill leonardo-office (chạy lại bash install.sh)')


CONFIG = json.load(open(_shared('config.json'), encoding='utf-8'))
LOGO = _shared(CONFIG['logo'])                                                                      # logo gốc nguyên bộ (hươu trên, chữ dưới): chỉ dùng ở trang đầu
COMPANY = CONFIG.get('company_word', CONFIG['company'])                                            # đầu trang, chân trang Word: tên tiếng Anh + mã số thuế, không ghi địa chỉ
WORDMARK = _shared(CONFIG['logo_wordmark'])                                                         # chỉ chữ LEONARDO: đầu trang các trang sau. Không ghép lại hươu và chữ theo cách khác logo gốc
P = dict(green='2F5D4A', soft='EEF3EF', grid='4A5A52', bad='8F2F1D', inp='FBF5E0', grey='6B6B6B')   # cùng bảng màu F0 với Excel
BODY_FONT = HEAD_FONT = 'Tahoma'                                                                    # một phông cho cả văn bản, cùng phông với Excel
BODY_PT, TABLE_PT, SMALL_PT = 10.5, 9.5, 8
MARGIN = dict(top=25, bottom=20, left=30, right=20)                                                # mm
TEXT_W = 210 - MARGIN['left'] - MARGIN['right']                                                    # 160 mm
INDENT = 15                                                                                        # mm: cột số mục; chữ mọi cấp thẳng hàng sau cột này
rgb = lambda h: RGBColor.from_string(h)

# Thứ tự phần tử con theo lược đồ OOXML. Word có thể từ chối mở file nếu các phần tử này sai thứ tự,
# nên mọi phần tử tự thêm đều đi qua _put() thay vì append().
ORDER = {
    'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd',
              'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription'],
    'tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
    'pPr': ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
            'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid',
            'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment', 'textboxTightWrap',
            'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'],
    'sectPr': ['footnotePr', 'endnotePr', 'type', 'pgSz', 'pgMar', 'paperSrc', 'pgBorders', 'lnNumType', 'pgNumType', 'cols', 'formProt', 'vAlign', 'noEndnote',
               'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings', 'sectPrChange'],
    'trPr': ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden'],
}


def _put(parent, child):
    """Đặt `child` vào `parent` đúng thứ tự lược đồ; thay phần tử cùng tên nếu đã có."""
    local = lambda e: e.tag.split('}')[1]
    order = ORDER[local(parent)]; name = local(child)
    old = parent.find(child.tag)
    if old is not None: parent.remove(old)
    k = order.index(name)
    for sib in list(parent):
        if local(sib) in order and order.index(local(sib)) > k:
            sib.addprevious(child); return child
    parent.append(child); return child


def _font(run, name=None, size=None, bold=None, italic=None, color=None):
    f = run.font
    if name:
        f.name = name; rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = rgb(color)
    return run


def _style_font(style, name, size, bold=False, italic=False, color=None):
    f = style.font; f.name = name; f.size = Pt(size); f.bold = bold; f.italic = italic
    if color: f.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]


def _shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); _put(tcPr, sh)


def _borders(el_pr, tag, spec):
    """spec: {'top': (sz, color) | None, ...}; sz tính theo 1/8 pt."""
    b = OxmlElement(tag)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if side not in spec: continue
        e = OxmlElement(f'w:{side}'); v = spec[side]
        if v is None: e.set(qn('w:val'), 'nil')
        else: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(v[0])); e.set(qn('w:space'), '0'); e.set(qn('w:color'), v[1])
        b.append(e)
    _put(el_pr, b)


def _ppr_border(pPr, side, sz, color, space=4):
    pb = OxmlElement('w:pBdr'); e = OxmlElement(f'w:{side}')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color); pb.append(e); _put(pPr, pb)


def _para_border(par, side, sz, color, space=4): _ppr_border(par._p.get_or_add_pPr(), side, sz, color, space)


def _val(tag, val):
    e = OxmlElement(tag); e.set(qn('w:val'), str(val)); return e


def _num_pr(pPr, num_id, level=0):
    """Gắn đánh số tự động của Word vào một đoạn hoặc một kiểu; num_id=0 là tắt đánh số."""
    np = OxmlElement('w:numPr'); np.append(_val('w:ilvl', level)); np.append(_val('w:numId', num_id)); _put(pPr, np)


def no_number(par):
    """Tắt số tự động cho một tiêu đề (dùng khi số đã nằm sẵn trong chữ, ví dụ lúc trình bày lại văn bản có sẵn)."""
    _num_pr(par._p.get_or_add_pPr(), 0); par.paragraph_format.left_indent = Mm(0); par.paragraph_format.first_line_indent = Mm(0)


def _fld(par, kind=None, code=None, **fmt):
    """Một mảnh của trường Word: kind = 'begin' | 'separate' | 'end', hoặc code = lệnh của trường."""
    r = par.add_run(); _font(r, **fmt)
    if kind: fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
    else: it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code; r._r.append(it)


def _field(par, code, **fmt):
    r = par.add_run(); _font(r, **fmt)
    for kind, text in (('begin', None), (None, code), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = text; r._r.append(it)


def _no_borders(table):
    _borders(table._tbl.tblPr, 'w:tblBorders', {s: None for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')})


def _widths(table, widths_mm, pad=True):
    """Khoá độ rộng bảng và từng cột (Word và các trình xem khác chỉ tôn trọng khi bảng ở chế độ cố định)."""
    tw = lambda mm: str(int(round(mm * 56.693)))
    table.autofit = False; pr = table._tbl.tblPr
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), tw(sum(widths_mm))); w.set(qn('w:type'), 'dxa'); _put(pr, w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); _put(pr, lay)
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), '100' if pad else '0'); e.set(qn('w:type'), 'dxa'); mar.append(e)
    _put(pr, mar)
    grid = table._tbl.tblGrid
    for gc, mm in zip(grid.findall(qn('w:gridCol')), widths_mm): gc.set(qn('w:w'), tw(mm))
    for row in table.rows:
        for c, mm in zip(row.cells, widths_mm): c.width = Mm(mm)


class Doc:
    """Một văn bản theo chuẩn Leonardo.

    Văn bản ngắn: title_block → [toc] → h1/h2/h3/p/bullets/numbered/info/table/note → signers → save.
    Văn bản dài:  cover → [toc] → start_body → h1/… → signers → save.
    """

    def __init__(self, header_title='', code='', letterhead=True):
        """letterhead=False: không đặt logo và thông tin Leonardo (dùng khi trình bày lại văn bản của tổ chức khác)."""
        self.letterhead = letterhead
        self.d = d = Document(); self.n = [0, 0, 0]; self.ntable = 0
        self.heads = []; self.indent = 0; self._toc_par = None; self._cover = False     # indent: 0 cho tới tiêu đề đầu tiên, sau đó thân bài lùi vào INDENT
        s = d.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297); s.orientation = WD_ORIENT.PORTRAIT
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = (Mm(MARGIN[k]) for k in ('top', 'bottom', 'left', 'right'))
        s.header_distance, s.footer_distance = Mm(10), Mm(10)
        st = d.styles
        _style_font(st['Normal'], BODY_FONT, BODY_PT); pf = st['Normal'].paragraph_format
        pf.line_spacing = 1.3; pf.space_after = Pt(6); pf.space_before = Pt(0)
        self._numbering()
        # tiêu đề màu đen; thứ bậc thể hiện bằng cỡ chữ, độ đậm và một đường kẻ mảnh phía trên tiêu đề cấp 1
        for lvl, (name, size, kw, before, after) in enumerate((('Heading 1', 12, dict(bold=True, color='000000'), 16, 6), ('Heading 2', 11, dict(bold=True, color='000000'), 10, 4),
                                                              ('Heading 3', BODY_PT, dict(bold=True, italic=True, color='000000'), 8, 3))):
            _style_font(st[name], HEAD_FONT, size, **kw); f = st[name].paragraph_format
            f.space_before, f.space_after, f.keep_with_next, f.line_spacing = Pt(before), Pt(after), True, 1.15
            _num_pr(st[name].element.get_or_add_pPr(), self.head_num, lvl)
        _ppr_border(st['Heading 1'].element.get_or_add_pPr(), 'top', 6, P['green'], space=6)
        for k, bold, left in ((1, True, 0), (2, False, INDENT)):                      # kiểu của mục lục
            ts = st.add_style(f'toc {k}', WD_STYLE_TYPE.PARAGRAPH); ts.base_style = st['Normal']; ts.font.bold = bold; f = ts.paragraph_format
            f.left_indent, f.first_line_indent, f.space_after, f.alignment = Mm(left + INDENT), Mm(-INDENT), Pt(3), WD_ALIGN_PARAGRAPH.LEFT
            f.tab_stops.add_tab_stop(Mm(left + INDENT)); f.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        self._header_footer(s, header_title, code)

    # ----- đánh số tự động của Word: tiêu đề 1. / 1.1 / 1.1.1 gắn với Heading 1–3; danh sách (a) / (i)
    def _numbering(self):
        n = self.d.part.numbering_part.element
        ids = lambda tag, attr: [int(e.get(qn(attr))) for e in n.findall(qn(tag))]
        self._next_abs = max(ids('w:abstractNum', 'w:abstractNumId'), default=0) + 1; self._next_num = max(ids('w:num', 'w:numId'), default=0) + 1
        self.head_num = self._new_num([('decimal', '%1.', INDENT, INDENT, 'Heading1'), ('decimal', '%1.%2', INDENT, INDENT, 'Heading2'), ('decimal', '%1.%2.%3', INDENT, INDENT, 'Heading3')])

    def _new_num(self, levels):
        """Thêm một định nghĩa đánh số nhiều cấp và trả về số hiệu để gắn vào đoạn. levels: [(kiểu số, mẫu chữ, lề trái mm, treo mm, kiểu đoạn gắn kèm)]."""
        n = self.d.part.numbering_part.element; tw = lambda mm: str(int(round(mm * 56.693)))
        aid, nid = self._next_abs, self._next_num; self._next_abs += 1; self._next_num += 1
        a = OxmlElement('w:abstractNum'); a.set(qn('w:abstractNumId'), str(aid)); a.append(_val('w:multiLevelType', 'multilevel'))
        for i, (fmt, text, left, hang, style) in enumerate(levels):
            l = OxmlElement('w:lvl'); l.set(qn('w:ilvl'), str(i)); l.append(_val('w:start', 1)); l.append(_val('w:numFmt', fmt))
            if style: l.append(_val('w:pStyle', style))
            l.append(_val('w:lvlText', text)); l.append(_val('w:lvlJc', 'left'))
            ppr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind'); ind.set(qn('w:left'), tw(left)); ind.set(qn('w:hanging'), tw(hang)); ppr.append(ind); l.append(ppr); a.append(l)
        first = n.find(qn('w:num'))                                                  # mọi abstractNum phải đứng trước num
        first.addprevious(a) if first is not None else n.append(a)
        e = OxmlElement('w:num'); e.set(qn('w:numId'), str(nid)); e.append(_val('w:abstractNumId', aid)); n.append(e); return nid

    # ----- đầu trang, chân trang
    def _footer(self, f, kind):
        """kind: 'full' = Trang x / tổng số trang; 'section' = Trang x / số trang của phần; 'roman' = chỉ số trang (i, ii) cho phần đầu; 'empty'."""
        f.is_linked_to_previous = False; fp = f.paragraphs[0]
        for c in list(fp._p): fp._p.remove(c)
        fp.paragraph_format.space_after = Pt(0)
        if kind == 'empty': return
        _para_border(fp, 'top', 4, P['grid'], space=6); g = dict(name=HEAD_FONT, size=SMALL_PT, color=P['grey'])
        if self.letterhead and kind != 'roman':
            fp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT); _font(fp.add_run(list(COMPANY.values())[0] + '\t'), **g)
        else: fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if kind == 'roman': _field(fp, 'PAGE', **g); return
        _font(fp.add_run('Trang '), **g); _field(fp, 'PAGE', **g); _font(fp.add_run(' / '), **g); _field(fp, 'NUMPAGES' if kind == 'full' else 'SECTIONPAGES', **g)

    def _header_footer(self, s, title, code):
        items = list(COMPANY.items())
        if not self.letterhead:                  # không nhận diện Leonardo: chỉ tên văn bản ở đầu trang và số trang ở chân trang
            hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; hp.paragraph_format.space_after = Pt(0)
            _font(hp.add_run(' · '.join(x for x in (title, code) if x)), HEAD_FONT, SMALL_PT, color=P['grey']); _para_border(hp, 'bottom', 6, P['green'], space=3)
            self._footer(s.footer, 'full'); return
        s.different_first_page_header_footer = True
        # trang đầu: logo gốc nguyên bộ cao 11 mm (bằng khối ba dòng thông tin bên phải), hai khối canh đáy, vạch xanh mảnh 0,75 pt
        h = s.first_page_header; h.is_linked_to_previous = False
        t = h.add_table(1, 2, Mm(TEXT_W)); _borders(t._tbl.tblPr, 'w:tblBorders', {'top': None, 'left': None, 'right': None, 'insideH': None, 'insideV': None, 'bottom': (6, P['green'])}); _widths(t, [40, TEXT_W - 40], pad=False)
        c0, c1 = t.rows[0].cells; c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(4); p0.paragraph_format.line_spacing = 1.1; p0.add_run().add_picture(LOGO, height=Mm(11))
        p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
        _font(p.add_run(items[0][1]), HEAD_FONT, 10, bold=True, color=P['green'])
        for k, v in items[1:]:
            q = c1.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; q.paragraph_format.space_after = Pt(0); q.paragraph_format.line_spacing = 1.1
            _font(q.add_run(f'{k} {v}'), HEAD_FONT, SMALL_PT, color=P['grey'])
        c1.paragraphs[-1].paragraph_format.space_after = Pt(5); h.add_paragraph().paragraph_format.space_after = Pt(0)
        # các trang sau: chỉ chữ LEONARDO rộng 30 mm + tên văn bản canh đáy; không kẻ đường (tránh trùng với đường kẻ của tiêu đề cấp 1)
        h2 = s.header; h2.is_linked_to_previous = False
        t2 = h2.add_table(1, 2, Mm(TEXT_W)); _no_borders(t2); _widths(t2, [50, TEXT_W - 50], pad=False)
        a, b = t2.rows[0].cells; a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        pa = a.paragraphs[0]; pa.paragraph_format.space_after = Pt(0); pa.add_run().add_picture(WORDMARK, width=Mm(30))
        q = b.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; q.paragraph_format.space_after = Pt(0)
        _font(q.add_run(' · '.join(x for x in (title, code) if x)), HEAD_FONT, SMALL_PT, color=P['grey'])
        h2.add_paragraph().paragraph_format.space_after = Pt(6)
        for f in (s.footer, s.first_page_footer): self._footer(f, 'full')

    # ----- trang bìa, mục lục, bắt đầu thân bài (cho văn bản dài)
    def cover(self, doc_type, title, subtitle=None, facts=None):
        """Trang bìa: loại văn bản, tên văn bản cỡ lớn canh trái, mô tả, và các dòng nhận diện (mã, phiên bản, ngày…) sau một vạch dọc."""
        d = self.d; self._cover = True
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(130); p.paragraph_format.space_after = Pt(6)
        if doc_type: _font(p.add_run(doc_type.upper()), HEAD_FONT, 11, bold=True, color=P['grey'])
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing = 1.1; p.paragraph_format.space_after = Pt(10)
        _font(p.add_run(title), HEAD_FONT, 24, bold=True, color=P['green'])
        if subtitle:
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; _font(p.add_run(subtitle), BODY_FONT, 11)
        d.add_paragraph().paragraph_format.space_after = Pt(30)
        if facts: self.info({k: (v or '…') for k, v in facts.items()}, label_w=34, bold=True)
        s = d.sections[0]; s.different_first_page_header_footer = True
        self._footer(s.first_page_footer, 'empty'); self._footer(s.footer, 'roman')

    def toc(self, title='Mục lục'):
        """Mục lục hai cấp, là trường TOC của Word: số trang do Word điền khi mở file (Word sẽ hỏi có cập nhật trường không)."""
        if self._cover: self.page_break()
        p = self.d.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(8); p.paragraph_format.keep_with_next = True
        _para_border(p, 'top', 6, P['green'], space=6); _font(p.add_run(title), HEAD_FONT, 14, bold=True)
        self._toc_par = self.d.add_paragraph()
        st = self.d.settings.element; uf = _val('w:updateFields', 'true')
        tail = next((c for c in st if c.tag.split('}')[1] in ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr')), None)
        tail.addprevious(uf) if tail is not None else st.append(uf)

    def _fill_toc(self):
        ents = [(l, n, t) for l, n, t in self.heads if l <= 2] or [(1, '', '')]; par = self._toc_par
        for i, (l, n, t) in enumerate(ents):
            if i:
                e = OxmlElement('w:p'); par._p.addnext(e); par = Paragraph(e, par._parent)
            par.style = self.d.styles[f'toc {l}']
            if i == 0: _fld(par, 'begin'); _fld(par, code=' TOC \\o "1-2" \\h \\z \\u '); _fld(par, 'separate')
            par.add_run(f'{n}\t{t}')
            if i == len(ents) - 1: _fld(par, 'end')
        if not self._cover: par.paragraph_format.space_after = Pt(14)

    def start_body(self):
        """Sau bìa và mục lục: sang phần mới, số trang đánh lại từ 1 (phần đầu đánh i, ii)."""
        s2 = self.d.add_section(WD_SECTION.NEW_PAGE); s1 = self.d.sections[0]
        for s, fmt in ((s1, 'lowerRoman'), (s2, 'decimal')):
            e = OxmlElement('w:pgNumType'); e.set(qn('w:fmt'), fmt); e.set(qn('w:start'), '1'); _put(s._sectPr, e)
        s2.different_first_page_header_footer = False; self._footer(s2.footer, 'section')

    # ----- khối tên văn bản và bảng kiểm soát tài liệu
    def title_block(self, doc_type, title, subtitle=None, control=None):
        if doc_type:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
            _font(p.add_run(doc_type.upper()), HEAD_FONT, 11, bold=True, color=P['grey'])
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
        _font(p.add_run(title.upper()), HEAD_FONT, 16, bold=True, color=P['green'])
        if subtitle:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _font(p.add_run(subtitle), BODY_FONT, BODY_PT, italic=True)
        if control: self.info(control, cols=2, label_w=32, bold=True)

    def info(self, items, cols=1, label_w=40, bold=False):
        """Bảng "nhãn – giá trị" không kẻ lưới: nhãn nhỏ màu xám bên trái, một vạch dọc xanh, giá trị bên phải. cols=2 xếp thành hai cặp cột."""
        items = list(items.items()) if isinstance(items, dict) else [tuple(x) for x in items]
        rows = (len(items) + cols - 1) // cols; pw = TEXT_W / cols
        ind = self.indent if cols == 1 else 0                                        # bảng một cặp cột lùi vào thẳng hàng với đoạn văn
        t = self.d.add_table(rows=rows, cols=2 * cols); _no_borders(t); _widths(t, [label_w, pw - label_w - ind] * cols)
        if ind: ti = OxmlElement('w:tblInd'); ti.set(qn('w:w'), str(int(round(ind * 56.693)))); ti.set(qn('w:type'), 'dxa'); _put(t._tbl.tblPr, ti)
        for i, (k, v) in enumerate(items):
            r, c = i % rows, 2 * (i // rows); a, b = t.cell(r, c), t.cell(r, c + 1)
            self._cell(a, k, color=P['grey'], size=SMALL_PT + 0.5); self._cell(b, v, bold=bold, size=BODY_PT)
            _borders(a._element.get_or_add_tcPr(), 'w:tcBorders', {'right': (6, P['green'])})
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def _cell(self, cell, text, bold=False, fill=None, align='l', color=None, italic=False, size=TABLE_PT):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; p = cell.paragraphs[0]
        p.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2); p.paragraph_format.line_spacing = 1.1
        _font(p.add_run('' if text is None else str(text)), HEAD_FONT, size, bold=bold, italic=italic, color=color)
        if fill: _shade(cell, fill)

    # ----- tiêu đề mục: số do Word tự đánh (1. / 1.1 / 1.1.1), nằm trong cột số rộng INDENT
    def _h(self, level, text):
        self.n[level - 1] += 1
        for i in range(level, 3): self.n[i] = 0
        num = '.'.join(str(x) for x in self.n[:level]) + ('.' if level == 1 else '')
        text = text.upper() if level == 1 else text
        self.heads.append((level, num, text)); self.indent = INDENT            # số chỉ lưu để dựng sẵn mục lục; trong văn bản không gõ số
        return self.d.add_heading(text, level=level)

    def h1(self, text): self._h(1, text)
    def h2(self, text): self._h(2, text)
    def h3(self, text): self._h(3, text)

    def p(self, text, align='justify', bold=False, italic=False):
        par = self.d.add_paragraph(); par.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
        if self.indent and align != 'center': par.paragraph_format.left_indent = Mm(self.indent)
        self._rich(par, text, bold, italic); return par

    def _rich(self, par, text, bold=False, italic=False):
        """**đậm** trong chuỗi được in đậm."""
        for i, part in enumerate(str(text).split('**')):
            if part: _font(par.add_run(part), BODY_FONT, BODY_PT, bold=bold or i % 2 == 1, italic=italic)

    def bullets(self, items, numbered=False):
        """Gạch đầu dòng, hoặc danh sách (a), (b)… do Word tự đánh; phần tử là một danh sách con thì thành cấp (i), (ii)…"""
        nid = self._new_num([('lowerLetter', '(%1)', self.indent + 9, 9, None), ('lowerRoman', '(%2)', self.indent + 18, 9, None)]) if numbered else None

        def one(it, lvl):
            par = self.d.add_paragraph(); f = par.paragraph_format; f.space_after = Pt(3); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if numbered:
                _num_pr(par._p.get_or_add_pPr(), nid, lvl); f.left_indent = Mm(self.indent + 9 * (lvl + 1)); f.first_line_indent = Mm(-9)
            else:
                f.left_indent = Mm(self.indent + 4.5); f.first_line_indent = Mm(-4.5); _font(par.add_run('–  '), BODY_FONT, BODY_PT)
            self._rich(par, it)
        for it in items:
            if isinstance(it, list):
                for sub in it: one(sub, 1)
            else: one(it, 0)

    def table(self, columns, rows, caption=None, widths=None, align=None, total_row=False):
        """columns: tiêu đề cột; rows: danh sách dòng; align: 'l'/'c'/'r' cho từng cột (mặc định: số canh phải)."""
        if caption:
            self.ntable += 1; c = self.d.add_paragraph(); c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(3); c.paragraph_format.keep_with_next = True
            _font(c.add_run(f'Bảng {self.ntable}. '), HEAD_FONT, TABLE_PT, bold=True, color=P['green']); _font(c.add_run(caption), HEAD_FONT, TABLE_PT)
        t = self.d.add_table(rows=1 + len(rows), cols=len(columns)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t._tbl.tblPr, 'w:tblBorders', {s: (4, P['grid']) for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')})
        if not widths: widths = [TEXT_W / len(columns)] * len(columns)
        _widths(t, [w * TEXT_W / sum(widths) for w in widths])
        isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
        def show(v):
            if isinstance(v, bool) or v is None: return '' if v is None else str(v)
            if isinstance(v, int): return f'{v:,}'.replace(',', '.') if v >= 0 else '(' + f'{-v:,}'.replace(',', '.') + ')'
            if isinstance(v, float):
                s = f'{abs(v):,.1f}'.replace(',', '§').replace('.', ',').replace('§', '.'); return s if v >= 0 else f'({s})'
            return str(v)
        for j, h in enumerate(columns): self._cell(t.cell(0, j), h, bold=True, fill=P['green'], align='c', color='FFFFFF')
        trp = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); _put(trp, th)   # lặp lại tiêu đề khi bảng sang trang
        for i, row in enumerate(rows, 1):
            last = total_row and i == len(rows)
            for j, v in enumerate(row):
                a = (align[j] if align else None) or ('r' if isnum(v) else 'l')
                self._cell(t.cell(i, j), show(v), bold=last, fill=P['soft'] if last else None, align=a)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def note(self, text, label='Lưu ý'):
        """Khung ghi chú: vạch xanh bên trái, nền nhạt."""
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t._tbl.tblPr, 'w:tblBorders', {'top': None, 'bottom': None, 'right': None, 'left': (24, P['green'])}); _widths(t, [TEXT_W])
        c = t.cell(0, 0); _shade(c, P['soft']); par = c.paragraphs[0]; par.paragraph_format.space_before = Pt(4); par.paragraph_format.space_after = Pt(4)
        _font(par.add_run(f'{label}: '), BODY_FONT, BODY_PT - 1, bold=True, color=P['green']); _font(par.add_run(text), BODY_FONT, BODY_PT - 1)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def signers(self, people, place_date=None):
        """people: [(chức danh, họ tên), ...] tối đa 3 cột; place_date: 'TP. Hồ Chí Minh, ngày … tháng … năm …'."""
        if place_date:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before = Pt(14); _font(p.add_run(place_date), BODY_FONT, BODY_PT, italic=True)
        t = self.d.add_table(rows=3, cols=len(people)); _no_borders(t); _widths(t, [TEXT_W / len(people)] * len(people))
        for j, (role, name) in enumerate(people):
            self._cell(t.cell(0, j), role.upper(), bold=True, align='c', size=11); self._cell(t.cell(1, j), '(Ký, ghi rõ họ tên)', italic=True, align='c', color=P['grey'], size=9)
            c = t.cell(2, j); c.paragraphs[0].paragraph_format.space_before = Pt(56); self._cell(c, name or '', bold=True, align='c', size=11)
            c.paragraphs[0].paragraph_format.space_before = Pt(56)
        for r in t.rows:
            trp = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); _put(trp, cs)

    def save(self, path):
        if self._toc_par is not None: self._fill_toc(); self._toc_par = None
        self.d.save(path)
