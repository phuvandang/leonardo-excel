"""Bộ chuẩn trình bày văn bản Word của Leonardo: lề, phông, hệ tiêu đề, bảng, đầu trang/chân trang, khối ký.

Dùng chung bảng màu, logo và thông tin công ty với skill leonardo-excel. Mọi hằng số trình bày nằm ở đây.
"""
import json, os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _shared(name):
    """Tìm file dùng chung: trong chính skill này trước, rồi tới skill leonardo-excel đặt cạnh."""
    for base in (ROOT, os.path.join(os.path.dirname(ROOT), 'leonardo-excel')):
        p = os.path.join(base, name)
        if os.path.exists(p): return p
    raise FileNotFoundError(f'Không tìm thấy {name}: cần cài cùng skill leonardo-excel (bash install.sh)')


CONFIG = json.load(open(_shared('config.json'), encoding='utf-8'))
LOGO = _shared(CONFIG['logo'])
P = dict(green='2F5D4A', soft='EEF3EF', grid='4A5A52', bad='8F2F1D', inp='FBF5E0', grey='6B6B6B')   # cùng bảng màu F0 với Excel
BODY_FONT, HEAD_FONT = 'Times New Roman', 'Tahoma'
BODY_PT, TABLE_PT, SMALL_PT = 12, 10, 8
MARGIN = dict(top=25, bottom=20, left=30, right=20)                                                # mm
TEXT_W = 210 - MARGIN['left'] - MARGIN['right']                                                    # 160 mm
rgb = lambda h: RGBColor.from_string(h)

# Thứ tự phần tử con theo lược đồ OOXML. Word có thể từ chối mở file nếu các phần tử này sai thứ tự,
# nên mọi phần tử tự thêm đều đi qua _put() thay vì append().
ORDER = {
    'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd',
              'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription'],
    'tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
    'pPr': ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
            'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid',
            'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment', 'textboxTightWrap',
            'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'],
    'trPr': ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden'],
}


def _put(parent, child):
    """Đặt `child` vào `parent` đúng thứ tự lược đồ; thay phần tử cùng tên nếu đã có."""
    local = lambda e: e.tag.split('}')[1]
    order = ORDER[local(parent)]; name = local(child)
    old = parent.find(child.tag)
    if old is not None: parent.remove(old)
    k = order.index(name)
    for sib in list(parent):
        if local(sib) in order and order.index(local(sib)) > k:
            sib.addprevious(child); return child
    parent.append(child); return child


def _font(run, name=None, size=None, bold=None, italic=None, color=None):
    f = run.font
    if name:
        f.name = name; rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = rgb(color)
    return run


def _style_font(style, name, size, bold=False, italic=False, color=None):
    f = style.font; f.name = name; f.size = Pt(size); f.bold = bold; f.italic = italic
    if color: f.color.rgb = rgb(color)
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]


def _shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); _put(tcPr, sh)


def _borders(el_pr, tag, spec):
    """spec: {'top': (sz, color) | None, ...}; sz tính theo 1/8 pt."""
    b = OxmlElement(tag)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if side not in spec: continue
        e = OxmlElement(f'w:{side}'); v = spec[side]
        if v is None: e.set(qn('w:val'), 'nil')
        else: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(v[0])); e.set(qn('w:space'), '0'); e.set(qn('w:color'), v[1])
        b.append(e)
    _put(el_pr, b)


def _para_border(par, side, sz, color, space=4):
    pPr = par._p.get_or_add_pPr(); pb = OxmlElement('w:pBdr'); e = OxmlElement(f'w:{side}')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color); pb.append(e); _put(pPr, pb)


def _field(par, code, **fmt):
    r = par.add_run(); _font(r, **fmt)
    for kind, text in (('begin', None), (None, code), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = text; r._r.append(it)


def _no_borders(table):
    _borders(table._tbl.tblPr, 'w:tblBorders', {s: None for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')})


def _widths(table, widths_mm, pad=True):
    """Khoá độ rộng bảng và từng cột (Word và các trình xem khác chỉ tôn trọng khi bảng ở chế độ cố định)."""
    tw = lambda mm: str(int(round(mm * 56.693)))
    table.autofit = False; pr = table._tbl.tblPr
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), tw(sum(widths_mm))); w.set(qn('w:type'), 'dxa'); _put(pr, w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); _put(pr, lay)
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), '100' if pad else '0'); e.set(qn('w:type'), 'dxa'); mar.append(e)
    _put(pr, mar)
    grid = table._tbl.tblGrid
    for gc, mm in zip(grid.findall(qn('w:gridCol')), widths_mm): gc.set(qn('w:w'), tw(mm))
    for row in table.rows:
        for c, mm in zip(row.cells, widths_mm): c.width = Mm(mm)


class Doc:
    """Một văn bản theo chuẩn Leonardo. Gọi lần lượt: title_block → h1/h2/h3/p/bullets/numbered/table/note → signers → save."""

    def __init__(self, header_title='', code=''):
        self.d = d = Document(); self.n = [0, 0, 0]; self.ntable = 0
        s = d.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297); s.orientation = WD_ORIENT.PORTRAIT
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = (Mm(MARGIN[k]) for k in ('top', 'bottom', 'left', 'right'))
        s.header_distance, s.footer_distance = Mm(10), Mm(10)
        st = d.styles
        _style_font(st['Normal'], BODY_FONT, BODY_PT); pf = st['Normal'].paragraph_format
        pf.line_spacing = 1.3; pf.space_after = Pt(6); pf.space_before = Pt(0)
        for name, size, kw, before, after in (('Heading 1', 13, dict(bold=True, color=P['green']), 18, 6), ('Heading 2', 12, dict(bold=True, color='000000'), 12, 4),
                                             ('Heading 3', 11, dict(bold=True, italic=True, color='000000'), 8, 3)):
            _style_font(st[name], HEAD_FONT, size, **kw); f = st[name].paragraph_format
            f.space_before, f.space_after, f.keep_with_next, f.line_spacing = Pt(before), Pt(after), True, 1.15
        self._header_footer(s, header_title, code)

    # ----- đầu trang, chân trang
    def _header_footer(self, s, title, code):
        s.different_first_page_header_footer = True
        # trang đầu: tiêu đề thư có logo và thông tin công ty
        h = s.first_page_header; h.is_linked_to_previous = False
        t = h.add_table(1, 2, Mm(TEXT_W)); _borders(t._tbl.tblPr, 'w:tblBorders', {'top': None, 'left': None, 'right': None, 'insideH': None, 'insideV': None, 'bottom': (12, P['green'])}); _widths(t, [45, TEXT_W - 45], pad=False)
        c0, c1 = t.rows[0].cells; c0.paragraphs[0].add_run().add_picture(LOGO, width=Mm(34))
        items = list(CONFIG['company'].items()); p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0)
        _font(p.add_run(items[0][1]), HEAD_FONT, 10, bold=True, color=P['green'])
        for k, v in items[1:]:
            q = c1.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; q.paragraph_format.space_after = Pt(0); q.paragraph_format.line_spacing = 1.1
            _font(q.add_run(f'{k} {v}'), HEAD_FONT, SMALL_PT, color=P['grey'])
        c1.paragraphs[-1].paragraph_format.space_after = Pt(4); h.add_paragraph().paragraph_format.space_after = Pt(0)
        # các trang sau: logo nhỏ + tên văn bản
        h2 = s.header; h2.is_linked_to_previous = False
        t2 = h2.add_table(1, 2, Mm(TEXT_W)); _borders(t2._tbl.tblPr, 'w:tblBorders', {'top': None, 'left': None, 'right': None, 'insideH': None, 'insideV': None, 'bottom': (6, P['green'])}); _widths(t2, [35, TEXT_W - 35], pad=False)
        a, b = t2.rows[0].cells; a.paragraphs[0].add_run().add_picture(LOGO, width=Mm(20)); b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        q = b.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; q.paragraph_format.space_after = Pt(0)
        _font(q.add_run(' · '.join(x for x in (title, code) if x)), HEAD_FONT, SMALL_PT, color=P['grey'])
        q.paragraph_format.space_after = Pt(3); h2.add_paragraph().paragraph_format.space_after = Pt(0)
        for f in (s.footer, s.first_page_footer):
            f.is_linked_to_previous = False; fp = f.paragraphs[0]; fp.paragraph_format.space_after = Pt(0)
            _para_border(fp, 'top', 4, P['grid'], space=6)
            fp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
            _font(fp.add_run(items[0][1] + '\tTrang '), HEAD_FONT, SMALL_PT, color=P['grey'])
            _field(fp, 'PAGE', name=HEAD_FONT, size=SMALL_PT, color=P['grey']); _font(fp.add_run(' / '), HEAD_FONT, SMALL_PT, color=P['grey'])
            _field(fp, 'NUMPAGES', name=HEAD_FONT, size=SMALL_PT, color=P['grey'])

    # ----- khối tên văn bản và bảng kiểm soát tài liệu
    def title_block(self, doc_type, title, subtitle=None, control=None):
        if doc_type:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
            _font(p.add_run(doc_type.upper()), HEAD_FONT, 11, bold=True, color=P['grey'])
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
        _font(p.add_run(title.upper()), HEAD_FONT, 16, bold=True, color=P['green'])
        if subtitle:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _font(p.add_run(subtitle), BODY_FONT, BODY_PT, italic=True)
        if control:
            items = list(control.items()); half = (len(items) + 1) // 2
            t = self.d.add_table(rows=half, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            _borders(t._tbl.tblPr, 'w:tblBorders', {s: (4, P['grid']) for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')})
            _widths(t, [34, 46, 34, 46])
            for i, (k, v) in enumerate(items):
                r, c = (i, 0) if i < half else (i - half, 2)
                self._cell(t.cell(r, c), k, bold=True, fill=P['soft']); self._cell(t.cell(r, c + 1), v)
            self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def _cell(self, cell, text, bold=False, fill=None, align='l', color=None, italic=False, size=TABLE_PT):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; p = cell.paragraphs[0]
        p.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2); p.paragraph_format.line_spacing = 1.1
        _font(p.add_run('' if text is None else str(text)), HEAD_FONT, size, bold=bold, italic=italic, color=color)
        if fill: _shade(cell, fill)

    # ----- tiêu đề mục, đánh số tự động 1. / 1.1 / 1.1.1
    def _h(self, level, text):
        self.n[level - 1] += 1
        for i in range(level, 3): self.n[i] = 0
        num = '.'.join(str(x) for x in self.n[:level]) + ('.' if level == 1 else '')
        self.d.add_heading(f'{num} {text.upper() if level == 1 else text}', level=level)

    def h1(self, text): self._h(1, text)
    def h2(self, text): self._h(2, text)
    def h3(self, text): self._h(3, text)

    def p(self, text, align='justify', bold=False, italic=False):
        par = self.d.add_paragraph(); par.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
        self._rich(par, text, bold, italic); return par

    def _rich(self, par, text, bold=False, italic=False):
        """**đậm** trong chuỗi được in đậm."""
        for i, part in enumerate(str(text).split('**')):
            if part: _font(par.add_run(part), BODY_FONT, BODY_PT, bold=bold or i % 2 == 1, italic=italic)

    def bullets(self, items, numbered=False):
        for k, it in enumerate(items, 1):
            par = self.d.add_paragraph(); f = par.paragraph_format; f.left_indent = Mm(11 if numbered else 10); f.first_line_indent = Mm(-5.5 if numbered else -4.5); f.space_after = Pt(3)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _font(par.add_run((f'{chr(96 + k)}) ' if numbered else '–  ')), BODY_FONT, BODY_PT); self._rich(par, it)

    def table(self, columns, rows, caption=None, widths=None, align=None, total_row=False):
        """columns: tiêu đề cột; rows: danh sách dòng; align: 'l'/'c'/'r' cho từng cột (mặc định: số canh phải)."""
        if caption:
            self.ntable += 1; c = self.d.add_paragraph(); c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(3); c.paragraph_format.keep_with_next = True
            _font(c.add_run(f'Bảng {self.ntable}. '), HEAD_FONT, TABLE_PT, bold=True, color=P['green']); _font(c.add_run(caption), HEAD_FONT, TABLE_PT)
        t = self.d.add_table(rows=1 + len(rows), cols=len(columns)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t._tbl.tblPr, 'w:tblBorders', {s: (4, P['grid']) for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')})
        if not widths: widths = [TEXT_W / len(columns)] * len(columns)
        _widths(t, [w * TEXT_W / sum(widths) for w in widths])
        isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
        def show(v):
            if isinstance(v, bool) or v is None: return '' if v is None else str(v)
            if isinstance(v, int): return f'{v:,}'.replace(',', '.') if v >= 0 else '(' + f'{-v:,}'.replace(',', '.') + ')'
            if isinstance(v, float):
                s = f'{abs(v):,.1f}'.replace(',', '§').replace('.', ',').replace('§', '.'); return s if v >= 0 else f'({s})'
            return str(v)
        for j, h in enumerate(columns): self._cell(t.cell(0, j), h, bold=True, fill=P['green'], align='c', color='FFFFFF')
        trp = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); _put(trp, th)   # lặp lại tiêu đề khi bảng sang trang
        for i, row in enumerate(rows, 1):
            last = total_row and i == len(rows)
            for j, v in enumerate(row):
                a = (align[j] if align else None) or ('r' if isnum(v) else 'l')
                self._cell(t.cell(i, j), show(v), bold=last, fill=P['soft'] if last else None, align=a)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def note(self, text, label='Lưu ý'):
        """Khung ghi chú: vạch xanh bên trái, nền nhạt."""
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t._tbl.tblPr, 'w:tblBorders', {'top': None, 'bottom': None, 'right': None, 'left': (24, P['green'])}); _widths(t, [TEXT_W])
        c = t.cell(0, 0); _shade(c, P['soft']); par = c.paragraphs[0]; par.paragraph_format.space_before = Pt(4); par.paragraph_format.space_after = Pt(4)
        _font(par.add_run(f'{label}: '), BODY_FONT, BODY_PT - 1, bold=True, color=P['green']); _font(par.add_run(text), BODY_FONT, BODY_PT - 1)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def signers(self, people, place_date=None):
        """people: [(chức danh, họ tên), ...] tối đa 3 cột; place_date: 'TP. Hồ Chí Minh, ngày … tháng … năm …'."""
        if place_date:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before = Pt(14); _font(p.add_run(place_date), BODY_FONT, BODY_PT, italic=True)
        t = self.d.add_table(rows=3, cols=len(people)); _no_borders(t); _widths(t, [TEXT_W / len(people)] * len(people))
        for j, (role, name) in enumerate(people):
            self._cell(t.cell(0, j), role.upper(), bold=True, align='c', size=11); self._cell(t.cell(1, j), '(Ký, ghi rõ họ tên)', italic=True, align='c', color=P['grey'], size=9)
            c = t.cell(2, j); c.paragraphs[0].paragraph_format.space_before = Pt(56); self._cell(c, name or '', bold=True, align='c', size=11)
            c.paragraphs[0].paragraph_format.space_before = Pt(56)
        for r in t.rows:
            trp = r._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); _put(trp, cs)

    def save(self, path): self.d.save(path)
