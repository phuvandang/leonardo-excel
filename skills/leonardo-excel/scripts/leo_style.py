"""Bộ chuẩn trình bày Excel của Leonardo: màu, phông, định dạng số và lớp Sheet dựng bảng.

Dùng chung cho new_report.py (tạo file mới) và restyle_build.py (trình bày lại file có sẵn).
Mọi hằng số trình bày nằm ở đây; không gõ lại mã màu ở nơi khác.
"""
import json, math, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG = json.load(open(os.path.join(ROOT, 'config.json'), encoding='utf-8'))
LOGO = os.path.join(ROOT, CONFIG['logo'])

# Bảng màu F0 "xanh rừng đậm"
P = dict(title='2F5D4A', head='2F5D4A', head_text='FFFFFF', group='EEF3EF', grid='4A5A52', good='2F5D4A', bad='8F2F1D', inp='FBF5E0', outside='A6A6A6')
CHART = ['2F5D4A', '6F9A85', 'C5CEC9', 'D8C48C', 'E6D8AC']
FONT = 'Tahoma'
NUM = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
NUM2 = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
PCT = '0.0%;(0.0%);"-"'
FORMATS = {'number': NUM, 'decimal': NUM2, 'percent': PCT, 'date': 'dd/mm/yyyy', 'general': 'General'}

F = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
_g = Side(style='thin', color=P['grid']); BOX = Border(left=_g, right=_g, top=_g, bottom=_g)
fill = lambda c: PatternFill('solid', fgColor=c)
Cn = Alignment(horizontal='center', vertical='center', wrap_text=True)
Rt = Alignment(horizontal='right', vertical='center')
Lf = Alignment(horizontal='left', vertical='center', wrap_text=True)
ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X', 'XI', 'XII', 'XIII', 'XIV', 'XV', 'XVI', 'XVII', 'XVIII', 'XIX', 'XX',
         'XXI', 'XXII', 'XXIII', 'XXIV', 'XXV', 'XXVI', 'XXVII', 'XXVIII', 'XXIX', 'XXX']
rom = lambda i: ROMAN[i] if i < len(ROMAN) else str(i + 1)


def nlines(text, width):
    """Ước lượng số dòng khi chữ xuống dòng trong cột rộng `width` (đơn vị độ rộng cột Excel)."""
    cap = max(1, int((7 * width + 5 - 6) / 6.5)); n = 0
    for para in str(text).split('\n'):
        cur = 0; n += 1
        for w in para.split(' '):
            L = len(w)
            if cur and cur + 1 + L > cap: n += 1; cur = 0
            while L > cap: n += 1; L -= cap
            cur = cur + (1 if cur else 0) + L
    return n


H = lambda n: max(20, math.ceil(13.2 * n + 5))


class Sheet:
    """Một sheet theo chuẩn: dòng 1-3 tên báo cáo/kỳ/đơn vị, dòng 4-9 thông tin công ty, từ dòng 11 là các khối bảng.

    Cột A = STT, B = nội dung, C = mã chỉ tiêu, từ D là cột giá trị.
    """

    def __init__(self, wb, name, title, widths, prefix='', subtitle=None, unit=None, period=None, report_date=None, info=True):
        self.ws = ws = wb.create_sheet(name); self.name = name; self.n = len(widths); self.last = CL(self.n)
        self.prefix = prefix; self.code = 0; self.widths = widths; self.nsec = 0; self.links = {}
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1): ws.column_dimensions[CL(i)].width = w
        L = self.L = CL(min(self.n, 12))
        ws.merge_cells(f'A1:{L}1'); self.put('A1', title, F(bold=True, color=P['title'], size=13), Cn); ws.row_dimensions[1].height = 30
        ws.merge_cells(f'A2:{L}2'); self.put('A2', subtitle, F(italic=True), Cn)
        ws.merge_cells(f'A3:{L}3'); self.put('A3', f"Đơn vị: {unit or CONFIG['unit']}", F(italic=True), Cn)
        if os.path.exists(LOGO):
            img = Image(LOGO); img.width, img.height = 113, 58; ws.add_image(img, 'A1')
        self.r = 5
        if info:
            rows = list(CONFIG['company'].items()) + [('Kỳ báo cáo:', period), ('Ngày lập báo cáo:', report_date), ('Người lập / Người duyệt:', CONFIG['signers'])]
            for i, (a, b) in enumerate(rows, start=4):
                ws.merge_cells(f'A{i}:B{i}'); ws.merge_cells(f'C{i}:{L}{i}'); self.put(f'A{i}', a); self.put(f'C{i}', b); ws.row_dimensions[i].height = 17
            self.r = 4 + len(rows) + 1
        self.limit = self.n

    def put(self, ref, v, font=None, al=None, fmt=None, fl=None, box=False):
        c = self.ws[ref]; c.value = v; c.font = font or F(); c.alignment = al or Alignment(vertical='center')
        if fmt: c.number_format = fmt
        if fl: c.fill = fill(fl)
        if box: c.border = BOX

    def sec(self, a, b):
        """Dòng tiêu đề mục: 'A.' / 'I' ở cột A, tên mục in hoa ở cột B."""
        self.put(f'A{self.r}', a, F(bold=True), Alignment(horizontal='center', vertical='center')); self.put(f'B{self.r}', b, F(bold=True))
        self.ws.row_dimensions[self.r].height = 17; self.r += 1

    def block(self, title, ncols=None):
        """Mở một khối bảng mới: cách một dòng, đánh số La Mã, giới hạn số cột giá trị của khối."""
        if self.nsec: self.gap()
        self.sec(rom(self.nsec), str(title).upper()); self.nsec += 1
        self.limit = self.n if ncols is None else min(self.n, 3 + max(ncols, 1))

    def header(self, labels, merge_from=None):
        r = self.r; n = 1
        for i in range(1, self.limit + 1):
            t = labels[i - 1] if i <= len(labels) else None
            self.put(f'{CL(i)}{r}', t, F(bold=True, color=P['head_text']), Cn, fl=P['head'], box=True)
            if t: n = max(n, nlines(t, self.widths[i - 1] * 0.9))
        if merge_from and self.limit > merge_from: self.ws.merge_cells(f'{CL(merge_from)}{r}:{CL(self.limit)}{r}')
        self.ws.row_dimensions[r].height = max(30, H(n)); self.r += 1

    def row(self, stt, label, vals=(), kind='item', merge_text=False, code=True):
        """Một dòng bảng. kind: item | group (đậm, nền nhạt) | sub (dòng tên nhóm, không mã).

        code: True = mã tự đánh số; False = để trống; (giá trị, ô nguồn) = lấy mã từ file gốc.
        vals: danh sách (giá trị, định dạng, ô nguồn). Có ô nguồn thì ô là công thức '=nguồn' và giá trị được ghi vào self.links
        để gắn làm giá trị đã tính khi ghép file. Giá trị là chuỗi bắt đầu bằng '=' được ghi làm công thức.
        Định dạng: None (chữ), 'input' (ô nhập nền vàng), 'status' (chữ đậm canh giữa) hoặc mã định dạng số.
        """
        r = self.r; bold = kind in ('group', 'sub'); fl = P['group'] if bold else None; f = F(bold=bold)
        c = None
        if isinstance(code, tuple):          # mã lấy từ ô nguồn (ví dụ mã nhân viên): (giá trị, ô nguồn)
            c = '=' + code[1]; self.links[f'C{r}'] = code[0]
        elif kind != 'sub' and code: self.code += 1; c = f'{self.prefix}{self.code:03d}'
        self.put(f'A{r}', stt, f, Cn, fl=fl, box=True); self.put(f'B{r}', label, f, Lf, fl=fl, box=True); self.put(f'C{r}', c, f, Cn, fl=fl, box=True)
        n = nlines(label or '', self.widths[1])
        for i in range(4, self.limit + 1):
            v, fmt, sref = (tuple(vals[i - 4]) + (None, None))[:3] if i - 4 < len(vals) else (None, None, None)
            ref = f'{CL(i)}{r}'
            if fmt == 'input': self.put(ref, v, F(), Lf, fl=P['inp'], box=True); continue
            if v is None: self.put(ref, None, f, Rt, None, fl, True); continue
            formula = isinstance(v, str) and v.startswith('=')
            text = isinstance(v, str) and not formula
            al = (Cn if fmt == 'status' else (Lf if (merge_text or len(v) > 24) else Cn)) if text else Rt
            font = F(bold=True) if fmt == 'status' else f
            nf = None if (text or fmt == 'status') else (fmt or NUM)
            if sref: self.put(ref, '=' + sref, font, al, nf, fl, True); self.links[ref] = v
            else: self.put(ref, v, font, Cn if (formula and fmt == 'status') else al, nf, fl, True)
            if text: n = max(n, nlines(v, sum(self.widths[3:self.limit]) if merge_text else self.widths[i - 1]))
        if merge_text and self.limit > 4: self.ws.merge_cells(f'D{r}:{CL(self.limit)}{r}')
        self.ws.row_dimensions[r].height = min(H(n), 400); self.r += 1
        return r

    def note(self, text):
        r = self.r; self.ws.merge_cells(f'A{r}:{self.L}{r}'); self.put(f'A{r}', text, F(italic=True), Lf)
        self.ws.row_dimensions[r].height = H(nlines(text, sum(self.widths[:min(self.n, 12)]))); self.r += 1

    def gap(self): self.r += 1

    def status_cf(self, rng):
        """Tô đỏ ô trạng thái có chữ 'vượt' hoặc 'lỗ', xanh cho ô còn lại."""
        a = rng.split(':')[0]
        self.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("vượt",{a})),ISNUMBER(SEARCH("lỗ",{a})))'], font=Font(color=P['bad'], bold=True)))
        self.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEN({a})>0'], font=Font(color=P['good'], bold=True)))

    def finish(self):
        """Nền trắng trong vùng báo cáo, xám bên ngoài; vùng in A4 ngang vừa một trang rộng."""
        ws = self.ws; ws.print_area = f'A1:{self.last}{self.r}'; ws.sheet_view.zoomScale = 100
        GREY = fill(P['outside']); WHITE = fill('FFFFFF')
        for r in range(1, self.r + 1):
            if ws.row_dimensions[r].height is None: ws.row_dimensions[r].height = 15
            for c in range(1, self.n + 1):
                cell = ws.cell(r, c)
                try:
                    if cell.fill.fill_type is None: cell.fill = WHITE
                except AttributeError: pass
        for i in range(1, self.n + 1): ws.column_dimensions[CL(i)].fill = GREY
        cd = ColumnDimension(ws, index=CL(self.n + 1), min=self.n + 1, max=16384); cd.fill = GREY; ws.column_dimensions[CL(self.n + 1)] = cd
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True); ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_options.horizontalCentered = True
