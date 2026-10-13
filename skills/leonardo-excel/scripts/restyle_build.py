"""Luồng "trình bày lại", bước 1: dựng các sheet trình bày theo chuẩn Leonardo từ một file Excel có sẵn.

Mỗi dòng nguồn có nội dung thành một dòng trình bày; mỗi ô số là công thức trỏ về đúng ô nguồn, không gõ lại số nào.
Kết quả là một file tạm (W) cùng hai file phụ .values.json và .meta.json, dùng cho restyle_merge.py.

Ví dụ:
  python restyle_build.py goc.xlsx W.xlsx \
      --sheet "BC DÒNG TIỀN::02 Dòng tiền::BÁO CÁO QUẢN TRỊ — DÒNG TIỀN::Tháng 08 năm 2026" \
      --period "Từ 01/08/2026 đến 31/08/2026"
"""
import argparse, datetime, json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as CL
from openpyxl.utils.datetime import to_excel
from leo_style import Sheet, NUM, NUM2, PCT, rom

ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
ap.add_argument('src'); ap.add_argument('out')
ap.add_argument('--sheet', action='append', required=True, help='"Tên sheet nguồn::Tên sheet mới::Tiêu đề::Dòng phụ" (lặp lại cho từng sheet)')
ap.add_argument('--unit'); ap.add_argument('--period'); ap.add_argument('--report-date')
ap.add_argument('--signers', help='"Người lập / Người duyệt" — chỉ truyền khi file gốc ghi rõ hoặc người dùng đã cho biết; không truyền thì để trống')
ap.add_argument('--label-width', type=float, default=44); ap.add_argument('--value-width', type=float, default=18.5)
ap.add_argument('--no-info', action='store_true', help='bỏ khối thông tin công ty')
ap.add_argument('--label-col', help='chữ cột nguồn dùng làm cột nội dung (mặc định: tự nhận). Ví dụ C cho cột Họ tên')
ap.add_argument('--code-col', help='chữ cột nguồn đưa vào cột thứ ba thay cho mã tự đánh số (ví dụ B cho mã nhân viên, hoặc cột Phòng ban khi bảng không có mã)')
ap.add_argument('--code-title', help='tiêu đề cột mã (mặc định: tiêu đề của --code-col trong file gốc, hoặc "Mã chỉ tiêu")')
A = ap.parse_args()
S = openpyxl.load_workbook(A.src, read_only=True, data_only=True)
clean = lambda t: ' '.join(str(t).split()) if isinstance(t, str) else t
isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
col_index = lambda L: openpyxl.utils.column_index_from_string(L.upper())
isdate = lambda v: isinstance(v, (datetime.datetime, datetime.date, datetime.time))


def load(name):
    rows = {}
    for row in S[name].iter_rows():
        d = {}; rn = None
        for c in row:
            v = c.value
            if v is None or (isinstance(v, str) and not v.strip()): continue
            bold = bool(c.font and c.font.b)
            filled = bool(c.fill and c.fill.fill_type and str(c.fill.fgColor.rgb)[-6:] not in ('FFFFFF', '000000'))
            d[c.column] = (v, bold, filled, c.number_format or 'General'); rn = c.row
        if d: rows[rn] = d
    return rows


def hdr_text(v):
    if isdate(v): return v.strftime('%d/%m/%Y')
    if isnum(v): return str(int(v)) if float(v).is_integer() else str(v)
    return clean(v)


def fmt_for(v, nf):
    if isdate(v): return 'dd/mm/yyyy'
    if '%' in nf: return PCT
    small = isinstance(v, float) and not v.is_integer() and abs(v) < 1000
    if nf == 'General': return 'General' if small else NUM
    return NUM2 if re.search(r'0\.0', nf) and small else NUM


def regions(rows):
    """Tách sheet thành các vùng bảng đặt cạnh nhau: một cột sau khoảng trống chỉ mở vùng mới nếu nó có cột nhãn riêng."""
    used = sorted({c for d in rows.values() for c in d}); regs = [[used[0]]]
    for c in used[1:]:
        nstr = sum(1 for d in rows.values() if c in d and isinstance(d[c][0], str))
        if c == regs[-1][-1] + 1 or nstr < 4: regs[-1].append(c)
        else: regs.append([c])
    return regs


def label_col(rows, cols):
    if A.label_col and col_index(A.label_col) in cols: return col_index(A.label_col)
    cnt = {c: sum(1 for d in rows.values() if c in d and isinstance(d[c][0], str)) for c in cols[:2]}
    return max(cols[:2], key=lambda c: (cnt[c], -c))


def render(sh, name, rows, cols, skipped):
    q = "'" + name.replace("'", "''") + "'"
    lab = label_col(rows, cols); before = [c for c in cols if c < lab]; cc = col_index(A.code_col) if A.code_col else None
    vcols = [c for c in cols if c > lab and c != cc]; items = []
    for rn in sorted(rows):
        d = {c: rows[rn][c] for c in cols if c in rows[rn]}
        if not d: continue
        strs = [c for c in d if isinstance(d[c][0], str)]; nums = [c for c in d if isnum(d[c][0]) and c not in before]; dates = [c for c in d if isdate(d[c][0])]
        lc = lab if lab in d and isinstance(d[lab][0], str) else next((c for c in before if c in d and isinstance(d[c][0], str) and c != cc), None)
        bold = any(d[c][1] for c in ([lc] if lc else list(d)[:1])); filled = bool(lc and d[lc][2])
        if not nums and (len(strs) + len(dates) >= 3 or (len(strs) == 2 and lc and d[lc][1])): cls = 'HEAD'
        elif not nums and not dates and len(strs) == 1 and lc and len(d) == 1: cls = 'TEXT1'
        else: cls = 'DATA'
        items.append(dict(rn=rn, d=d, cls=cls, lc=lc, bold=bold, filled=filled, nums=nums))
    for i, it in enumerate(items):          # dòng chữ đứng ngay trên dòng tiêu đề cột là tên khối
        if it['cls'] == 'TEXT1':
            for j in range(i + 1, min(i + 3, len(items))):
                if items[j]['cls'] == 'HEAD':
                    it['cls'] = 'TITLE'
                    for k in range(i + 1, j): items[k]['cls'] = 'NOTE'
                    break
                if items[j]['nums']: break
    cur = None
    for it in items:                         # số cột thật của từng khối
        if it['cls'] == 'HEAD': cur = it; it['ext'] = 0
        if cur is not None and it['cls'] in ('HEAD', 'DATA'):
            cur['ext'] = max(cur['ext'], max([vcols.index(c) + 1 for c in it['d'] if c in vcols] or [0]))
            if it['cls'] == 'DATA' and it['nums']: cur['hasnum'] = True
    blocks = []                              # bảng danh sách mà đa số dòng đều in đậm/tô nền: chỉ coi dòng có nhãn nằm ở cột phía trước là dòng tổng
    for it in items:
        if it['cls'] == 'HEAD': blocks.append([])
        elif it['cls'] == 'DATA' and blocks: blocks[-1].append(it)
    for blk in blocks:
        if blk and sum(1 for it in blk if it['bold'] or it['filled']) > len(blk) / 2:
            for it in blk:                   # dòng nhóm: nhãn nằm ở cột phía trước, có ký hiệu nhóm in đậm ở cột phía trước (A, B, I, II…), hoặc không có nhãn mà chữ đậm nằm ở cột số liệu
                mark = any(isinstance(it['d'][c][0], str) and it['d'][c][1] for c in before if c in it['d'] and c != cc)
                it['bold'] = it['filled'] = bool((it['lc'] and it['lc'] != lab) or mark or (it['lc'] is None and it['bold']))
                it['nostt'] = it['bold'] and not mark          # dòng tổng của bảng danh sách: không đánh số
    in_table = False; title = None; notes = []; gi = ii = 0; textblock = False
    for it in items:
        d, rn, cls = it['d'], it['rn'], it['cls']
        first_txt = next((clean(d[c][0]) for c in d if isinstance(d[c][0], str)), None)
        if cls == 'TITLE': title = first_txt; continue
        if cls == 'NOTE': notes.append(': '.join(hdr_text(d[c][0]) for c in sorted(d))); continue
        if cls == 'HEAD':
            hl = hdr_text(d[it['lc']][0]) if it['lc'] else 'Chỉ tiêu'
            textblock = it['ext'] == 1 and not it.get('hasnum')
            sh.block(title or hl, ncols=(min(sh.n, 9) - 3) if textblock else it['ext'])
            for t in notes: sh.note(t)
            sh.header(['STT', hl, A.code_title or (hdr_text(d[cc][0]) if cc and cc in d else 'Mã chỉ tiêu')] + [hdr_text(d[c][0]) if c in d else None for c in vcols], merge_from=4 if textblock else None)
            in_table = True; title = None; notes = []; gi = ii = 0; continue
        if not in_table and cls == 'TEXT1': title = first_txt; continue
        if not in_table:
            skipped.append((name, rn, [(CL(c), str(d[c][0])[:30]) for c in sorted(d)][:6])); continue
        label = clean(d[it['lc']][0]) if it['lc'] else None
        vals = []
        for c in vcols:
            if c not in d or c == it['lc']: vals.append((None, None, None)); continue
            v, _, _, nf = d[c]
            vals.append((to_excel(v) if isdate(v) else v, None if isinstance(v, str) else fmt_for(v, nf), f'{q}!{CL(c)}{rn}'))
        if cls == 'TEXT1' or not any(v is not None for v, _, _ in vals):
            sh.row('', label, [], 'sub'); ii = 0; continue
        present = [v for v, _, _ in vals if v is not None]
        merge = len(present) == 1 and isinstance(present[0], str) and vals[0][0] is not None and (textblock or len(present[0]) > 30)
        src = next((d[c][0] for c in before if c in d and c != cc and c != it['lc']), None)      # số thứ tự hoặc ký hiệu nhóm có sẵn trong bản gốc
        src = int(src) if isnum(src) and float(src).is_integer() else (clean(src) if isinstance(src, str) else None)
        if not A.label_col: src = None                    # chỉ bảng danh sách (có --label-col) mới giữ số thứ tự gốc; báo cáo thì đánh lại theo từng nhóm
        if it['bold'] or it['filled']: stt = '' if it.get('nostt') else (src if isinstance(src, str) else rom(gi)); gi += 1; ii = 0; kind = 'group'
        elif label is None and merge: stt = ''; kind = 'item'
        else: ii += 1; stt = src if src is not None else ii; kind = 'item'
        code = (d[cc][0], f'{q}!{CL(cc)}{rn}') if cc and cc in d and cc != it['lc'] else (not cc)
        sh.row(stt, label if label is not None else ('' if merge or kind == 'group' else '(không có tên trong bản gốc)'), vals, kind, merge_text=merge, code=code)


wb = Workbook(); wb.remove(wb.active); skipped = []; VAL = {}; meta = dict(sources=[], new=[]); summary = []
for k, spec in enumerate(A.sheet, 1):
    parts = [p.strip() for p in spec.split('::')]
    sname, new = parts[0], parts[1]; title = parts[2] if len(parts) > 2 else new.upper(); sub = parts[3] if len(parts) > 3 else None
    assert sname in S.sheetnames, f'Không có sheet nguồn: {sname}'; assert new not in S.sheetnames, f'Tên sheet mới trùng sheet có sẵn: {new}'
    rows = load(sname); regs = regions(rows)
    nv = [len([c for c in cols if c > label_col(rows, cols)]) for cols in regs]
    sh = Sheet(wb, new, title, [7, A.label_width, 11] + [A.value_width] * max(max(nv), 3), f'{k:02d}', sub, unit=A.unit, period=A.period, report_date=A.report_date, info=not A.no_info, signers=A.signers)
    sh.sec('A.', title.split('— ')[-1])
    for cols in regs:
        if len(cols) < 2 and not any(isnum(d[cols[0]][0]) for d in rows.values() if cols[0] in d):
            skipped += [(sname, rn, [(CL(cols[0]), str(d[cols[0]][0])[:30])]) for rn, d in sorted(rows.items()) if cols[0] in d]; continue
        render(sh, sname, rows, cols, skipped)
    sh.finish(); VAL.update({f'{new}|{ref}': v for ref, v in sh.links.items()})
    meta['sources'].append(sname); meta['new'].append(new); summary.append((new, sh.r, sh.n))
wb.save(A.out)
js = lambda v: v if isinstance(v, (int, float, str, bool)) or v is None else str(v)
json.dump({k: js(v) for k, v in VAL.items()}, open(A.out + '.values.json', 'w', encoding='utf-8'), ensure_ascii=False)
meta['skipped'] = skipped; json.dump(meta, open(A.out + '.meta.json', 'w', encoding='utf-8'), ensure_ascii=False)
if not A.no_info and not A.signers: print('⚠ Dòng "Người lập / Người duyệt" đang để trống: hỏi người dùng, không tự điền.')
print('Đã dựng:', summary, '| ô liên kết:', len(VAL), '| dòng nguồn không đưa vào:', len(skipped))
for s in skipped[:30]: print('  BỎ QUA', s[0], 'dòng', s[1], str(s[2])[:140])
